"""Fija lo que muestran hoy los listados del inicio y del panel (prueba de caracterización).

El resultado esperado está en listados_esperados.json. Si un cambio de comportamiento es intencional,
se regenera con:  REGENERAR_LISTADOS=1 python -m unittest tests.test_listados
"""
import importlib
import json
import os
import sys
import tempfile
import unittest
from unittest import mock

from flask import template_rendered

import core.database
from core.database import execute_query

ESPERADOS = os.path.join(os.path.dirname(__file__), "listados_esperados.json")

CONSULTAS_INICIO = [
    {}, {"sort": "price_desc"}, {"sort": "alpha_asc"}, {"sort": "alpha_desc"}, {"sort": "desconocido"},
    {"grado": "HG"}, {"grado": "Todos"}, {"q": "gundam"}, {"q": "  zaku "}, {"page": 2},
    {"per_page": 48}, {"per_page": 7},
]
CONSULTAS_PANEL = [
    {}, {"sort_admin": "price_asc"}, {"sort_admin": "price_desc"}, {"sort_admin": "alpha_asc"}, {"sort_admin": "alpha_desc"},
    {"grado_admin": "HG"}, {"q_admin": "gundam"}, {"page": 2}, {"per_page": 48},
]


def cargar_datos():
    """Casos que distinguen las reglas de cada listado."""
    q = lambda sql, args=(): execute_query(sql, args, commit=True)
    q("INSERT INTO tiendas (id, nombre) VALUES (1, 'A'), (2, 'B')")
    # (id, nombre, grado, es_visible, disponible)
    for p in [(1, "ZAKU HG", "HG", 1, 1), (2, "GUNDAM EXIA RG", "RG", 1, 1), (3, "SINANJU MG", "MG", 0, 1),
              (4, "NU GUNDAM HG", "HG", 1, 0), (5, "HARO SD", "SD", 1, 1), (6, "BARBATOS HG", "HG", 0, 0),
              (7, "AERIAL HG", "HG", 1, 1)]:
        q("INSERT INTO productos (id, nombre_estandar, grado, es_visible, disponible) VALUES (?, ?, ?, ?, ?)", p)
    # (pub id, producto, tienda, precio actual, activa, [(precio, fecha), ...])
    publicaciones = [
        (1, 1, 1, 15000, 1, [(16000, "2026-01-01 12:00:00"), (15000, "2026-01-05 12:00:00")]),
        (2, 1, 2, 14000, 1, [(14000, "2026-01-03 12:00:00")]),
        (3, 2, 1, 35000, 1, [(30000, "2026-01-01 12:00:00"), (35000, "2026-01-04 12:00:00")]),  # subió: inicio y panel difieren
        (4, 3, 1, 60000, 1, [(60000, "2026-01-02 12:00:00")]),
        (5, 4, 1, 20000, 0, [(20000, "2026-01-02 12:00:00")]),
        (6, 6, 2, 18000, 0, [(18000, "2026-01-06 12:00:00")]),
        (7, 7, 2, 18000, 1, [(18000, "2026-01-06 12:00:00")]),
    ]
    # Relleno para que haya más de una página
    for n in range(25):
        q("INSERT INTO productos (id, nombre_estandar, grado) VALUES (?, ?, 'HG')", (100 + n, f"RELLENO {n:02d} HG"))
        publicaciones.append((100 + n, 100 + n, 1, 50000 + n, 1, [(50000 + n, f"2026-02-{1 + n % 20:02d} 12:00:00")]))
    for pub_id, prod, tienda, precio, activa, historial in publicaciones:
        q("INSERT INTO publicacion_producto (id, producto_id, tienda_id, url, precio, es_url_activa) VALUES (?, ?, ?, ?, ?, ?)",
          (pub_id, prod, tienda, f"https://t.cl/{pub_id}", precio, activa))
        for precio_h, fecha in historial:
            q("INSERT INTO historial_precios (publicacion_producto_id, precio, fecha_registro) VALUES (?, ?, ?)", (pub_id, precio_h, fecha))


class TestListados(unittest.TestCase):
    def setUp(self):
        fd, self.db = tempfile.mkstemp(suffix=".db")
        os.close(fd)
        self.parche = mock.patch.object(core.database, "DB_NAME", self.db)
        self.parche.start()
        sys.modules.pop("app", None)
        self.app = importlib.import_module("app").app  # init_db() corre sobre la base temporal
        cargar_datos()
        self.cliente = self.app.test_client()
        with self.cliente.session_transaction() as s:
            s["admin_logged"] = True

    def tearDown(self):
        self.parche.stop()
        os.remove(self.db)

    def contexto(self, ruta, consulta, plantilla):
        """Variables con que la ruta renderizó la plantilla."""
        capturado = {}
        def guardar(sender, template, context, **extra):
            if template.name == plantilla:
                capturado.update(context)
        with template_rendered.connected_to(guardar, self.app):
            # follow_redirects: /producto/<id> redirige a la URL con nombre
            self.assertEqual(self.cliente.get(ruta, query_string=consulta, follow_redirects=True).status_code, 200)
        return capturado

    def resultado_actual(self):
        resultado = {"inicio": [], "panel": []}
        for consulta in CONSULTAS_INICIO:
            c = self.contexto("/", consulta, "public/inicio.html")
            resultado["inicio"].append({"consulta": consulta, "total_pages": c["total_pages"], "datos": [list(f) for f in c["datos"]]})
        for consulta in CONSULTAS_PANEL:
            c = self.contexto("/admin/panel", consulta, "admin/admin_panel.html")
            resultado["panel"].append({"consulta": consulta, "total_pages": c["total_pages"],
                                       "lista_productos": [list(f) for f in c["lista_productos"]]})
        # Ficha de producto: varias tiendas, precio que subió, publicación inactiva y producto sin publicaciones
        resultado["producto"] = []
        for id in (1, 2, 4, 5):
            c = self.contexto(f"/producto/{id}", {}, "public/producto.html")
            resultado["producto"].append({"id": id, "nombre": c["nombre_estandar"], "tiendas": c["tiendas"],
                                          "fechas": json.loads(c["fechas_json"]), "datasets": json.loads(c["datasets_json"])})
        resultado["producto_inexistente"] = self.cliente.get("/producto/999").status_code
        # Estadísticas del panel (los eventos solo por texto: la fecha cambia en cada corrida)
        c = self.contexto("/admin/panel", {}, "admin/admin_panel.html")
        resultado["estadisticas"] = {k: c[k] for k in ("usuario", "total_productos", "ultima_fecha", "escaneados_hoy", "alerta_activa", "mensaje_alerta")}
        resultado["estadisticas"]["eventos"] = [accion for _, accion in c["lista_logs"]]
        # Excel exportado
        import io
        from openpyxl import load_workbook
        respuesta = self.cliente.get("/admin/exportar_excel")
        columnas, *filas = load_workbook(io.BytesIO(respuesta.data))["Historial Global"].iter_rows(values_only=True)
        # Los precios son REAL en la base: 50019.0 se compara como "50019"
        texto = lambda v: str(int(v)) if isinstance(v, float) and v.is_integer() else str(v)
        resultado["excel"] = {"columnas": list(columnas), "filas": [[texto(v) for v in fila] for fila in filas]}
        return resultado

    def test_listados_no_cambian(self):
        actual = self.resultado_actual()
        if os.environ.get("REGENERAR_LISTADOS"):
            with open(ESPERADOS, "w", encoding="utf-8") as f:
                json.dump(actual, f, ensure_ascii=False, indent=1)
            self.skipTest("listados_esperados.json regenerado")
        with open(ESPERADOS, encoding="utf-8") as f:
            esperado = json.load(f)
        for seccion in ("inicio", "panel"):
            for e, a in zip(esperado[seccion], actual[seccion]):
                self.assertEqual(e, a, f"{seccion} con {e['consulta']}")
        for seccion in ("producto", "producto_inexistente", "estadisticas", "excel"):
            self.assertEqual(esperado[seccion], actual[seccion], seccion)


if __name__ == "__main__":
    unittest.main()
